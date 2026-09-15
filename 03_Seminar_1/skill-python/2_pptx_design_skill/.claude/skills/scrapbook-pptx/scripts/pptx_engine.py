"""Pressed Flower Scrapbook 스타일 PPTX 엔진.

reference-slides.html 의 디자인 토큰/컴포넌트를 python-pptx 로 옮긴다.
- 좌표는 레퍼런스와 같은 1280x720 px 기준 (1px = 9525 EMU)
- 글자, 도형, 찢어진 종이, 우표 카드는 PowerPoint 네이티브 도형 (편집 가능)
- 압화/고사리/아이콘/차트/다이어그램은 SVG 로 삽입 (PowerPoint 2016+ 는 벡터, 그 외는 PNG 대체 이미지)
- 같은 명세로 HTML 미리보기(preview.html)도 함께 만든다

사용 예:
    deck = Deck("output")
    s = deck.slide("kraft")
    with s.card("paper", 170, 110, 840, 490, rot=-1.2):
        s.text(90, 104, 700, 200, ["{A}I와 지역", "노동시장"], style="display", size=92)
    s.svg("sprig", 990, 290, 270, 560, rot=-10)
    deck.save("deck.pptx")
"""
from __future__ import annotations

import html
import math
import os
import random
import re
import shutil
import subprocess
import tempfile
from contextlib import contextmanager
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525
W, H = 1280, 720

# ---- 디자인 토큰 (reference-slides.html :root 와 동일) ----
C = {
    "brown900": "4A2410", "brown800": "5C2E14", "brown600": "7A4A2A",
    "kraft": "B39373", "kraft_light": "CDB295", "paper": "EAE0CC", "paper2": "F3EBDD",
    "peach": "F8E7D3", "ledger": "EFD6A0", "ink": "3B2414", "muted": "7B6552",
    "yellow": "E6BE4A", "lilac": "A98ABD", "sage": "6D7C53", "white": "FFFDF8",
    "up": "5E7A3A", "down": "A2452B",
    "tint": "E1D4BD",        # paper 위 brown 8% (시사점 박스, 표 줄무늬)
    "stitch": "8F6B55",      # 다크 카드 위 점선
    "rule": "CDBBA3",        # paper 위 점선 구분선
}
FONT = {"latin": "Arial", "display": "Arial Black", "ea": "맑은 고딕"}
CSS_FONT = "Arial,'Malgun Gothic','Apple SD Gothic Neo',sans-serif"
CSS_DISPLAY = "'Arial Black',Arial,'Malgun Gothic','Apple SD Gothic Neo',sans-serif"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"

# 텍스트 스타일 프리셋: (size px, bold, font, color, line-height, letter-spacing em)
STYLES = {
    "display": dict(size=54, bold=True, font="display", color="brown800", lh=1.06),
    "h2": dict(size=28, bold=True, color="brown800", lh=1.3),
    "h3": dict(size=21, bold=True, color="brown800", lh=1.35),
    "body": dict(size=19, color="ink", lh=1.65),
    "small": dict(size=16, color="ink", lh=1.6),
    "caption": dict(size=14, color="muted", lh=1.5),
    "kicker": dict(size=14, bold=True, color="brown600", lh=1.2, spacing=0.18),
}


# =====================================================================
# 텍스처 (PIL/numpy 로 생성)
# =====================================================================
def _save_jpg(arr, path):
    Image.fromarray(np.clip(arr, 0, 255).astype("uint8")).save(path, quality=88, dpi=(96, 96))


def _grain_tile(path, hex_color, sigma=6.0, size=256, seed=1):
    rng = np.random.default_rng(seed)
    base = np.array([int(hex_color[i:i + 2], 16) for i in (0, 2, 4)], dtype=float)
    _save_jpg(base + rng.normal(0, sigma, (size, size, 1)), path)


def _ledger_tile(path, seed=3):
    rng = np.random.default_rng(seed)
    img = np.array([239, 214, 160], dtype=float) * np.ones((180, 240, 3))
    line, red = np.array([92, 46, 20.0]), np.array([160, 60, 40.0])
    for y in range(17, 180, 18):
        img[y] = img[y] * 0.84 + line * 0.16
    for x in (118, 119, 238, 239):
        img[:, x] = img[:, x] * 0.78 + red * 0.22
    _save_jpg(img + rng.normal(0, 4, (180, 240, 1)), path)


def _kraft_texture(path, w=1920, h=1080, seed=7):
    """구겨진 크라프트지: ridged noise 높이맵 + 사광 조명."""
    rng = np.random.default_rng(seed)
    hmap = np.zeros((h, w))
    for s, wt in ((300, 1.0), (150, 0.7), (70, 0.4), (32, 0.2)):
        gw, gh = w // s + 3, h // s + 3
        small = Image.fromarray((rng.random((gh, gw)) * 255).astype("uint8"))
        n = np.asarray(small.resize((gw * s, gh * s), Image.BICUBIC), dtype=float)[:h, :w] / 255
        hmap += wt * (0.55 * (1 - np.abs(2 * n - 1)) + 0.45 * n)
    gy, gx = np.gradient(hmap)
    d = gx + gy
    shade = np.clip(0.5 + 0.12 * d / (d.std() + 1e-9), 0, 1)
    base = np.array([179, 147, 115], dtype=float)
    _save_jpg(base * (0.86 + 0.24 * shade[..., None]) + rng.normal(0, 4, (h, w, 1)), path)


# =====================================================================
# 가장자리 모양
# =====================================================================
def torn_points(w, h, seed, step=9, amp=3.2):
    """찢어진 종이 가장자리 폴리곤 (로컬 좌표)."""
    rnd = random.Random(seed)
    pts = []

    def edge(p0, p1, normal):
        length = math.dist(p0, p1)
        n = max(2, int(length / step))
        drift = 0.0
        for i in range(n):
            t = i / n
            drift = max(-amp, min(amp, drift + rnd.uniform(-1.6, 1.6)))
            off = drift + rnd.uniform(-amp * 0.6, amp * 0.6)
            pts.append((p0[0] + (p1[0] - p0[0]) * t + normal[0] * off,
                        p0[1] + (p1[1] - p0[1]) * t + normal[1] * off))

    edge((0, 0), (w, 0), (0, 1))
    edge((w, 0), (w, h), (-1, 0))
    edge((w, h), (0, h), (0, -1))
    edge((0, h), (0, 0), (1, 0))
    return pts


def stamp_points(w, h, pitch=16, r=5, seg=6):
    """우표 톱니(반원 구멍) 가장자리 폴리곤."""
    pts = []

    def edge(p0, p1, inward):
        length = math.dist(p0, p1)
        ux, uy = (p1[0] - p0[0]) / length, (p1[1] - p0[1]) / length
        pts.append(p0)
        k = int((length - pitch) // pitch)
        start = (length - k * pitch) / 2
        for i in range(k + 1):
            c = start + i * pitch
            for j in range(seg + 1):
                a = math.pi * j / seg
                along, depth = c - r * math.cos(a), r * math.sin(a)
                pts.append((p0[0] + ux * along + inward[0] * depth, p0[1] + uy * along + inward[1] * depth))

    edge((0, 0), (w, 0), (0, 1))
    edge((w, 0), (w, h), (-1, 0))
    edge((w, h), (0, h), (0, -1))
    edge((0, h), (0, 0), (1, 0))
    return pts


# =====================================================================
# SVG 그래픽 (PowerPoint SVG 렌더러 호환: <use>/필터/CSS 없이 평면 요소만)
# =====================================================================
def svg_doc(w, h, body):
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}" '
            f'font-family="Malgun Gothic, Arial, sans-serif">{body}</svg>')


def svg_flower(x, y, s, kind="y"):
    fill, stroke, center = (("#E6BE4A", "#C99A2E", "#A8652A") if kind == "y" else ("#B596C6", "#8E6FA3", "#F1DFA6"))
    k = s / 24
    return (f'<g transform="translate({x + s / 2:.1f} {y + s / 2:.1f}) scale({k:.3f})">'
            f'<g fill="{fill}" stroke="{stroke}" stroke-width="0.6">'
            '<ellipse cx="0" cy="-5.5" rx="4" ry="5.5"/><ellipse cx="0" cy="5.5" rx="4" ry="5.5"/>'
            '<ellipse cx="-5.5" cy="0" rx="5.5" ry="4"/><ellipse cx="5.5" cy="0" rx="5.5" ry="4"/></g>'
            f'<circle cx="0" cy="0" r="2.4" fill="{center}"/></g>')


def svg_sprig():
    stems = "".join(f'<path d="{d}" fill="none" stroke="#6D7C53" stroke-width="3" stroke-linecap="round"/>' for d in (
        "M104 420 C100 340 112 260 96 180 S74 70 66 24", "M98 214 C122 184 140 150 150 96",
        "M95 158 C72 136 52 112 40 72", "M101 292 C132 272 152 250 166 212", "M100 250 C78 236 58 222 44 190"))
    leaves = ('<path d="M103 372 q36 -34 74 -16 q-38 30 -74 16z" fill="#7E8B5E"/>'
              '<path d="M101 336 q-38 -30 -72 -8 q36 26 72 8z" fill="#6D7C53"/>'
              '<path d="M99 268 q-30 -14 -50 4 q28 12 50 -4z" fill="#7E8B5E"/>')
    flowers = "".join(svg_flower(x, y, s, k) for x, y, s, k in (
        (48, 4, 36, "y"), (76, 18, 28, "y"), (42, 30, 26, "l"), (20, 52, 34, "l"), (46, 64, 24, "y"), (16, 82, 24, "l"),
        (132, 72, 36, "y"), (158, 96, 26, "l"), (126, 102, 24, "y"), (150, 190, 34, "l"), (172, 210, 24, "y"),
        (24, 170, 32, "y"), (50, 188, 22, "l"), (84, 120, 22, "l")))
    return svg_doc(200, 420, stems + leaves + flowers)


def svg_fern():
    parts = ['<path d="M45 400 C44 300 50 160 44 6" stroke="#5F6E48" stroke-width="2.5" fill="none"/><g fill="#6D7C53">']
    for i, rx in enumerate((18, 18, 17, 16, 15, 13, 12, 10, 8, 6, 4)):
        y = 360 - i * 32
        ry = max(2.5, rx / 3)
        lx, rxc = 45 - rx * 0.85, 45 + rx * 0.85
        parts.append(f'<ellipse cx="{lx:.1f}" cy="{y}" rx="{rx}" ry="{ry:.1f}" transform="rotate(25 {lx:.1f} {y})"/>'
                     f'<ellipse cx="{rxc:.1f}" cy="{y - 6}" rx="{rx}" ry="{ry:.1f}" transform="rotate(-25 {rxc:.1f} {y - 6})"/>')
    parts.append("</g>")
    return svg_doc(90, 400, "".join(parts))


# =====================================================================
# 브라우저로 SVG -> PNG 대체 이미지
# =====================================================================
def find_browser():
    cands = [os.environ.get("PPTX_BROWSER", ""),
             r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
             r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
             r"C:\Program Files\Google\Chrome\Application\chrome.exe",
             "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
             shutil.which("google-chrome") or "", shutil.which("chromium") or "", shutil.which("msedge") or ""]
    return next((c for c in cands if c and Path(c).exists()), None)


class Deck:
    def __init__(self, out_dir):
        self.out = Path(out_dir)
        self.assets = self.out / "assets"
        self.assets.mkdir(parents=True, exist_ok=True)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(W * PX), Emu(H * PX)
        self.slides: list[Slide] = []
        self._svg_parts = {}
        self._svg_sizes = {}
        self._seed = 100
        self._browser = find_browser()
        self._profile = Path(tempfile.gettempdir()) / "pptx_engine_browser_profile"
        self._make_textures()
        self.add_svg("sprig", svg_sprig())
        self.add_svg("fern", svg_fern())

    # ---- 자원 ----
    def _make_textures(self):
        t = self.assets
        jobs = {
            "tex_paper.jpg": lambda p: _grain_tile(p, C["paper"], 6, seed=1),
            "tex_paper2.jpg": lambda p: _grain_tile(p, C["paper2"], 5, seed=2),
            "tex_peach.jpg": lambda p: _grain_tile(p, C["peach"], 5, seed=4),
            "tex_dark.jpg": lambda p: _grain_tile(p, C["brown800"], 7, seed=5),
            "tex_ledger.jpg": _ledger_tile,
            "tex_kraft.jpg": _kraft_texture,
        }
        for name, fn in jobs.items():
            if not (t / name).exists():
                fn(t / name)

    def add_svg(self, name, svg_text):
        """SVG 문자열을 assets/<name>.svg 로 저장하고 PNG 대체 이미지를 만든다."""
        svg_path, png_path = self.assets / f"{name}.svg", self.assets / f"{name}.png"
        m = re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg_text)
        w, h = float(m.group(1)), float(m.group(2))
        self._svg_sizes[name] = (w, h)
        changed = not svg_path.exists() or svg_path.read_text(encoding="utf-8") != svg_text
        svg_path.write_text(svg_text, encoding="utf-8")
        if changed or not png_path.exists():
            self._render_png(svg_path, png_path, w, h)
        return name

    def _render_png(self, svg_path, png_path, w, h):
        if not self._browser:
            Image.new("RGBA", (2, 2), (0, 0, 0, 0)).save(png_path)
            print(f"[경고] 브라우저를 찾지 못해 {png_path.name} 는 투명 대체 이미지로 저장됨")
            return
        page = svg_path.with_suffix(".render.html")
        page.write_text(f'<html><body style="margin:0;background:transparent;overflow:hidden">'
                        f'<img src="{svg_path.name}" style="display:block;width:{w}px;height:{h}px"></body></html>',
                        encoding="utf-8")
        subprocess.run([self._browser, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--no-first-run",
                        f"--user-data-dir={self._profile}", "--default-background-color=00000000",
                        "--force-device-scale-factor=2", f"--window-size={int(math.ceil(w))},{int(math.ceil(h))}",
                        f"--screenshot={png_path.resolve()}", page.resolve().as_uri()],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        page.unlink(missing_ok=True)

    def svg_part(self, name):
        if name not in self._svg_parts:
            blob = (self.assets / f"{name}.svg").read_bytes()
            self._svg_parts[name] = Part(PackURI(f"/ppt/media/svg_{name}.svg"), "image/svg+xml", self.prs.part.package, blob)
        return self._svg_parts[name]

    def next_seed(self):
        self._seed += 1
        return self._seed

    # ---- 슬라이드 ----
    def slide(self, bg="peach"):
        s = Slide(self, bg)
        self.slides.append(s)
        return s

    def save(self, pptx_name, preview_name="preview.html"):
        path = self.out / pptx_name
        self.prs.save(path)
        self._write_preview(self.out / preview_name)
        return path

    def _write_preview(self, path):
        css = (f"body{{margin:0;background:#D8CCBC;font-family:{CSS_FONT}}}"
               ".s{position:relative;width:1280px;height:720px;overflow:hidden;margin:24px auto;box-shadow:0 8px 24px rgba(0,0,0,.25)}"
               ".s *{box-sizing:border-box;word-break:keep-all}.a{position:absolute}.t p{margin:0}"
               "@page{size:1280px 720px;margin:0}@media print{body{background:none}.s{margin:0;box-shadow:none;page-break-after:always}}")
        body = "\n".join(f'<section class="s">{"".join(s.html)}</section>' for s in self.slides)
        path.write_text(f'<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>Preview</title>'
                        f"<style>{css}</style></head><body>{body}</body></html>", encoding="utf-8")


class Slide:
    def __init__(self, deck: Deck, bg):
        self.deck = deck
        self.sl = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.shapes = self.sl.shapes
        self.stack = []
        self.html = []
        self._pid = 0
        if bg == "kraft":
            self._fill_rect(0, 0, W, H, "tex_kraft.jpg", tile=False)
        else:
            self._fill_rect(0, 0, W, H, {"peach": "tex_peach.jpg", "brown": "tex_dark.jpg"}[bg])

    # ---- 좌표 변환 (card 안쪽 로컬 좌표 -> 슬라이드 좌표) ----
    def _xf(self, x, y, w, h):
        rot = 0.0
        for fx, fy, fw, fh, frot in reversed(self.stack):
            cx, cy = fx + x + w / 2, fy + y + h / 2
            px, py = fx + fw / 2, fy + fh / 2
            a = math.radians(frot)
            dx, dy = cx - px, cy - py
            cx, cy = px + dx * math.cos(a) - dy * math.sin(a), py + dx * math.sin(a) + dy * math.cos(a)
            x, y, rot = cx - w / 2, cy - h / 2, rot + frot
        return x, y, w, h, rot

    def _pt(self, x, y):
        gx, gy, _, _, _ = self._xf(x, y, 0, 0)
        return gx, gy

    @contextmanager
    def frame(self, x, y, w, h, rot=0.0):
        self.stack.append((x, y, w, h, rot))
        try:
            yield
        finally:
            self.stack.pop()

    @staticmethod
    def _box_css(x, y, w, h, rot, extra=""):
        t = f"transform:rotate({rot:.2f}deg);" if rot else ""
        return f'left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;{t}{extra}'

    # ---- 채움 도우미 ----
    def _set_blip_fill(self, shp, img_name, tile=True):
        _, rid = self.sl.part.get_or_add_image_part(str(self.deck.assets / img_name))
        sppr = shp._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
            for e in sppr.findall(qn(tag)):
                sppr.remove(e)
        mode = ('<a:tile tx="0" ty="0" sx="100000" sy="100000" flip="none" algn="tl"/>' if tile
                else "<a:stretch><a:fillRect/></a:stretch>")
        el = etree.fromstring(f'<a:blipFill xmlns:a="{A_NS}" xmlns:r="{R_NS}" rotWithShape="1">'
                              f'<a:blip r:embed="{rid}"/>{mode}</a:blipFill>')
        geom = sppr.find(qn("a:custGeom"))
        if geom is None:
            geom = sppr.find(qn("a:prstGeom"))
        geom.addnext(el)

    @staticmethod
    def _shadow(shp, blur=14, dist=7, alpha=30):
        sppr = shp._element.spPr
        sppr.append(etree.fromstring(
            f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="{blur * PX}" dist="{dist * PX}" dir="5400000" '
            f'algn="t" rotWithShape="0"><a:srgbClr val="281405"><a:alpha val="{alpha * 1000}"/></a:srgbClr>'
            f"</a:outerShdw></a:effectLst>"))

    def _fill_rect(self, x, y, w, h, img, tile=True):
        shp = self.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
        self._set_blip_fill(shp, img, tile)
        shp.line.fill.background()
        src = f"assets/{img}"
        size = "100% 100%" if not tile else "256px auto"
        self.html.append(f'<div class="a" style="{self._box_css(x, y, w, h, 0)}background:url({src});background-size:{size}"></div>')

    def _poly(self, pts, x, y, w, h, img, tile, shadow):
        gx, gy, _, _, rot = self._xf(x, y, w, h)
        fb = self.shapes.build_freeform(int((gx + pts[0][0]) * PX), int((gy + pts[0][1]) * PX), scale=1.0)
        fb.add_line_segments([(int((gx + px) * PX), int((gy + py) * PX)) for px, py in pts[1:]], close=True)
        shp = fb.convert_to_shape()
        self._set_blip_fill(shp, img, tile)
        shp.line.fill.background()
        if shadow:
            self._shadow(shp)
        shp.rotation = rot
        # preview
        self._pid += 1
        pid = f"p{id(self)}_{self._pid}"
        minx, miny = min(p[0] for p in pts), min(p[1] for p in pts)
        pw, ph = max(p[0] for p in pts) - minx, max(p[1] for p in pts) - miny
        poly = " ".join(f"{px - minx:.1f},{py - miny:.1f}" for px, py in pts)
        pat = (f'<pattern id="{pid}" patternUnits="userSpaceOnUse" width="256" height="256">'
               f'<image href="assets/{img}" width="256" height="256"/></pattern>') if tile else \
              (f'<pattern id="{pid}" patternUnits="userSpaceOnUse" width="{pw:.0f}" height="{ph:.0f}">'
               f'<image href="assets/{img}" width="{pw:.0f}" height="{ph:.0f}" preserveAspectRatio="xMidYMid slice"/></pattern>')
        if img == "tex_ledger.jpg":
            pat = pat.replace('width="256" height="256"><image', 'width="240" height="180"><image').replace(
                'width="256" height="256"/>', 'width="240" height="180"/>')
        filt = "filter:drop-shadow(0 7px 7px rgba(40,20,5,.3));" if shadow else ""
        self.html.append(f'<div class="a" style="{self._box_css(gx + minx, gy + miny, pw, ph, rot, filt)}">'
                         f'<svg width="{pw:.0f}" height="{ph:.0f}" style="display:block"><defs>{pat}</defs>'
                         f'<polygon points="{poly}" fill="url(#{pid})"/></svg></div>')
        return shp

    # ---- 컴포넌트 ----
    SCRAP_TEX = {"paper": "tex_paper.jpg", "paper2": "tex_paper2.jpg", "kraft": "tex_kraft.jpg",
                 "ledger": "tex_ledger.jpg", "dark": "tex_dark.jpg"}

    def scrap(self, kind, x, y, w, h, rot=0.0, shadow=None):
        """찢어진 종이 조각. kind: paper | paper2 | kraft | ledger | dark"""
        if shadow is None:
            shadow = kind in ("paper", "paper2")
        with self.frame(x, y, w, h, rot):
            self._poly(torn_points(w, h, self.deck.next_seed()), 0, 0, w, h, self.SCRAP_TEX[kind],
                       tile=kind != "kraft", shadow=shadow)

    @contextmanager
    def card(self, kind, x, y, w, h, rot=0.0, shadow=None):
        """종이 조각을 그리고, 그 안쪽을 로컬 좌표(0,0 = 카드 좌상단)로 사용."""
        self.scrap(kind, x, y, w, h, rot, shadow)
        with self.frame(x, y, w, h, rot):
            yield

    @contextmanager
    def stamp(self, x, y, w, h, rot=0.0, badge=None):
        """우표형 다크 카드 (톱니 가장자리 + 점선 스티치 + 번호 배지)."""
        with self.frame(x, y, w, h, rot):
            self._poly(stamp_points(w, h), 0, 0, w, h, "tex_dark.jpg", tile=True, shadow=True)
            self.rect(9, 9, w - 18, h - 18, fill=None, line="stitch", line_w=1.5, dash=True)
            if badge is not None:
                self.oval(w / 2 - 24, -24, 48, 48, str(badge), fill="paper2", line="brown800", line_w=3,
                          size=19, style="display", color="brown800")
            yield

    def rect(self, x, y, w, h, fill="paper", line=None, line_w=1.0, dash=False, radius=0, shadow=False,
             text=None, **tkw):
        gx, gy, gw, gh, rot = self._xf(x, y, w, h)
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shp = self.shapes.add_shape(kind, Emu(int(gx * PX)), Emu(int(gy * PX)), Emu(int(gw * PX)), Emu(int(gh * PX)))
        if radius:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        self._style_shape(shp, fill, line, line_w, dash)
        if shadow:
            self._shadow(shp, blur=12, dist=5, alpha=30)
        shp.rotation = rot
        if text is not None:
            self._fill_text(shp.text_frame, text, anchor="m", align=tkw.pop("align", "c"), **tkw)
        css = self._shape_css(fill, line, line_w, dash) + (f"border-radius:{radius}px;" if radius else "")
        if shadow:
            css += "box-shadow:0 5px 12px rgba(40,20,5,.3);"
        inner = self._text_html(text, anchor="m", align="c", **tkw) if text is not None else ""
        self.html.append(f'<div class="a" style="{self._box_css(gx, gy, gw, gh, rot, css)}">{inner}</div>')
        return shp

    def oval(self, x, y, w, h, text=None, fill="brown800", line=None, line_w=1.0, shadow=False, **tkw):
        gx, gy, gw, gh, rot = self._xf(x, y, w, h)
        shp = self.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(gx * PX)), Emu(int(gy * PX)), Emu(int(gw * PX)), Emu(int(gh * PX)))
        self._style_shape(shp, fill, line, line_w, False)
        if shadow:
            self._shadow(shp, blur=14, dist=6, alpha=30)
        shp.rotation = rot
        tkw.setdefault("style", "display")
        tkw.setdefault("color", "paper2")
        tkw.setdefault("size", 18)
        if text is not None:
            self._fill_text(shp.text_frame, text, anchor="m", align="c", **tkw)
        css = self._shape_css(fill, line, line_w, False) + "border-radius:50%;"
        if shadow:
            css += "box-shadow:0 6px 14px rgba(40,20,5,.3);"
        inner = self._text_html(text, anchor="m", align="c", **tkw) if text is not None else ""
        self.html.append(f'<div class="a" style="{self._box_css(gx, gy, gw, gh, rot, css)}">{inner}</div>')
        return shp

    def line(self, x1, y1, x2, y2, color="rule", width=2.0, dash=True):
        (ax, ay), (bx, by) = self._pt(x1, y1), self._pt(x2, y2)
        con = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(ax * PX)), Emu(int(ay * PX)),
                                        Emu(int(bx * PX)), Emu(int(by * PX)))
        con.line.color.rgb = RGBColor.from_string(C.get(color, color))
        con.line.width = Pt(width * 0.75)
        if dash:
            con.line.dash_style = MSO_LINE.DASH
        minx, miny = min(ax, bx) - 4, min(ay, by) - 4
        dasharr = 'stroke-dasharray="8 6"' if dash else ""
        self.html.append(f'<svg class="a" style="left:{minx:.1f}px;top:{miny:.1f}px;overflow:visible" width="1" height="1">'
                         f'<line x1="{ax - minx:.1f}" y1="{ay - miny:.1f}" x2="{bx - minx:.1f}" y2="{by - miny:.1f}" '
                         f'stroke="#{C.get(color, color)}" stroke-width="{width}" {dasharr}/></svg>')

    def svg(self, name, x, y, w, h, rot=0.0):
        """assets/<name>.svg 를 SVG(벡터) + PNG 대체 이미지로 삽입."""
        gx, gy, gw, gh, frot = self._xf(x, y, w, h)
        rot = rot + frot
        pic = self.shapes.add_picture(str(self.deck.assets / f"{name}.png"), Emu(int(gx * PX)), Emu(int(gy * PX)),
                                      Emu(int(gw * PX)), Emu(int(gh * PX)))
        pic.rotation = rot
        rid = self.sl.part.relate_to(self.deck.svg_part(name), RT.IMAGE)
        blip = pic._element.find(qn("p:blipFill")).find(qn("a:blip"))
        blip.append(etree.fromstring(
            f'<a:extLst xmlns:a="{A_NS}"><a:ext uri="{SVG_EXT_URI}">'
            f'<asvg:svgBlip xmlns:asvg="{ASVG_NS}" xmlns:r="{R_NS}" r:embed="{rid}"/></a:ext></a:extLst>'))
        self.html.append(f'<img class="a" src="assets/{name}.svg" style="{self._box_css(gx, gy, gw, gh, rot)}">')
        return pic

    # ---- 텍스트 ----
    def text(self, x, y, w, h, content, style="body", align="l", anchor="t", **kw):
        """content: 문자열 또는 문단 리스트. 인라인 표기: **굵게**, {섞기 글자}, ^^강조(굵은 갈색)^^"""
        gx, gy, gw, gh, rot = self._xf(x, y, w, h)
        tb = self.shapes.add_textbox(Emu(int(gx * PX)), Emu(int(gy * PX)), Emu(int(gw * PX)), Emu(int(gh * PX)))
        tb.rotation = rot
        self._fill_text(tb.text_frame, content, style=style, align=align, anchor=anchor, **kw)
        self.html.append(f'<div class="a t" style="{self._box_css(gx, gy, gw, gh, rot)}">'
                         f'{self._text_html(content, style=style, align=align, anchor=anchor, **kw)}</div>')
        return tb

    @staticmethod
    def _opts(style, kw):
        o = dict(size=19, bold=False, font="latin", color="ink", lh=1.5, spacing=0.0, space=0, mix_color=None)
        o.update(STYLES.get(style, {}))
        o.update({k: v for k, v in kw.items() if v is not None})
        if o["mix_color"] is None:
            o["mix_color"] = "ledger" if o["color"] in ("paper2", "white") else "brown600"
        return o

    @staticmethod
    def _segments(text):
        out = []
        for tok in re.split(r"(\*\*.+?\*\*|\{.+?\}|\^\^.+?\^\^)", text):
            if not tok:
                continue
            if tok.startswith("**"):
                out.append((tok[2:-2], "bold"))
            elif tok.startswith("^^"):
                out.append((tok[2:-2], "accent"))
            elif tok.startswith("{"):
                out.append((tok[1:-1], "mix"))
            else:
                out.append((tok, ""))
        return out

    def _fill_text(self, tf, content, style="body", align="l", anchor="t", **kw):
        o = self._opts(style, kw)
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        paras = content if isinstance(content, list) else [content]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = Pt(o["size"] * 0.75 * o["lh"])
            if o["space"] and i < len(paras) - 1:
                p.space_after = Pt(o["space"] * 0.75)
            for seg, mark in self._segments(ptxt):
                r = p.add_run()
                r.text = seg
                f = r.font
                size, bold, font, color = o["size"], o["bold"], o["font"], o["color"]
                if mark == "bold":
                    bold = True
                elif mark == "accent":
                    bold, color = True, ("ledger" if o["color"] in ("paper2", "white") else "brown800")
                elif mark == "mix":
                    size, font, bold, color = o["size"] * 0.9, "latin", True, o["mix_color"]
                f.size = Pt(size * 0.75)
                f.bold = bold
                f.name = FONT[font]
                f.color.rgb = RGBColor.from_string(C.get(color, color))
                rpr = r._r.get_or_add_rPr()
                latin = rpr.find(qn("a:latin"))
                ea = etree.SubElement(rpr, qn("a:ea"))
                ea.set("typeface", FONT["ea"])
                latin.addnext(ea)
                if o["spacing"]:
                    rpr.set("spc", str(int(o["spacing"] * size * 0.75 * 100)))

    def _text_html(self, content, style="body", align="l", anchor="t", **kw):
        o = self._opts(style, kw)
        paras = content if isinstance(content, list) else [content]
        fam = CSS_DISPLAY if o["font"] == "display" else CSS_FONT
        out = []
        for i, ptxt in enumerate(paras):
            spans = []
            for seg, mark in self._segments(ptxt):
                s = html.escape(seg)
                if mark == "bold":
                    s = f"<b>{s}</b>"
                elif mark == "accent":
                    col = C["ledger"] if o["color"] in ("paper2", "white") else C["brown800"]
                    s = f'<b style="color:#{col}">{s}</b>'
                elif mark == "mix":
                    s = (f'<span style="font-family:{CSS_FONT};font-weight:700;font-size:.9em;'
                         f'color:#{C.get(o["mix_color"], o["mix_color"])}">{s}</span>')
                spans.append(s)
            mb = f"margin-bottom:{o['space']}px;" if o["space"] and i < len(paras) - 1 else ""
            out.append(f'<p style="{mb}">{"".join(spans) or "&nbsp;"}</p>')
        just = {"t": "flex-start", "m": "center", "b": "flex-end"}[anchor]
        ls = f"letter-spacing:{o['spacing']}em;" if o["spacing"] else ""
        return (f'<div style="height:100%;display:flex;flex-direction:column;justify-content:{just};'
                f'font-family:{fam};font-size:{o["size"]}px;line-height:{o["lh"]};font-weight:{700 if o["bold"] else 400};'
                f'color:#{C.get(o["color"], o["color"])};text-align:{ {"l": "left", "c": "center", "r": "right"}[align]};{ls}">'
                f'{"".join(out)}</div>')

    # ---- 도형 스타일 ----
    @staticmethod
    def _style_shape(shp, fill, line, line_w, dash):
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(C.get(fill, fill))
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = RGBColor.from_string(C.get(line, line))
            shp.line.width = Pt(line_w * 0.75)
            if dash:
                shp.line.dash_style = MSO_LINE.DASH

    @staticmethod
    def _shape_css(fill, line, line_w, dash):
        css = f"background:#{C.get(fill, fill)};" if fill else ""
        if line:
            css += f"border:{line_w}px {'dashed' if dash else 'solid'} #{C.get(line, line)};"
        return css
